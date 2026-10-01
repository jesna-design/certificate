"""Private CSV/XLSX roster validation and import command."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import secrets
import hashlib
import sys
import tempfile
import uuid
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from portal_security import safe_certificate_url
from portal_storage import connect, data_dir, database_path


EMAIL_RE = re.compile(r"^[A-Z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Z0-9](?:[A-Z0-9.-]{0,251}[A-Z0-9])?\.[A-Z]{2,63}$", re.I)
ALIASES = {
    "name": {"name", "registered name", "participant name", "name of participant"},
    "email": {"email", "email id", "registered email", "registered email id", "e mail id"},
    "cert1": {
        "certificate 1", "certificate 1 link", "certificate 1 url", "sttp certificate",
        "sttp certifictae", "sttp cert", "completion certificate", "completion certificate link",
    },
    "cert2": {
        "certificate 2", "certificate 2 link", "certificate 2 url", "software certificate",
        "software certificate link", "software cert",
    },
}


def _normal_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip()


def _value(cell: Any) -> str:
    value = getattr(cell, "value", cell)
    if value is None:
        hyperlink = getattr(cell, "hyperlink", None)
        value = getattr(hyperlink, "target", "") if hyperlink else ""
    if value is None:
        return ""
    return str(value).strip()


def _row_mapping(headers: list[Any]) -> dict[str, int]:
    normalized = [_normal_header(value) for value in headers]
    mapping: dict[str, int] = {}
    for key, aliases in ALIASES.items():
        alias_set = {_normal_header(alias) for alias in aliases}
        for index, label in enumerate(normalized):
            if label in alias_set:
                mapping[key] = index
                break
    return mapping


def _load_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        sample = source.read(4096)
        source.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(source, dialect)
        rows = list(reader)
    if not rows:
        raise ValueError("The CSV file is empty")
    mapping = _row_mapping(rows[0])
    _require_identity_columns(mapping)
    return _records_from_rows(rows[1:], mapping, first_source_row=2)


def _load_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt to import Excel workbooks") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows()
            mapping = {}
            header_row_number = 0
            for row_number, row in enumerate(iterator, start=1):
                cells = list(row)
                values = [_value(cell) for cell in cells]
                candidate = _row_mapping(values)
                if "name" in candidate and "email" in candidate:
                    mapping = candidate
                    header_row_number = row_number
                    break
                if row_number >= 10:
                    break
            if not mapping:
                continue
            # Re-read the sheet because a header may be below a title or note row.
            source_rows = []
            for row_number, row in enumerate(sheet.iter_rows(min_row=header_row_number + 1), start=header_row_number + 1):
                source_rows.append((row_number, list(row)))
            return _records_from_rows(source_rows, mapping, first_source_row=header_row_number + 1)
    finally:
        workbook.close()
    raise ValueError("No worksheet has recognizable participant name and email columns")


def _require_identity_columns(mapping: dict[str, int]) -> None:
    missing = {"name", "email"} - set(mapping)
    if missing:
        raise ValueError("Could not identify required columns: " + ", ".join(sorted(missing)))


def _records_from_rows(rows: list[Any], mapping: dict[str, int], first_source_row: int) -> list[dict[str, Any]]:
    _require_identity_columns(mapping)
    records = []
    for offset, raw in enumerate(rows):
        if isinstance(raw, tuple) and len(raw) == 2 and isinstance(raw[0], int):
            source_row, cells = raw
            values = [_value(cell) for cell in cells]
        else:
            source_row, values = first_source_row + offset, [_value(cell) for cell in raw]
        get = lambda key: values[mapping[key]] if key in mapping and mapping[key] < len(values) else ""
        record = {
            "source_row": source_row,
            "name": get("name").strip(),
            "email_raw": get("email").strip(),
            "cert1_raw": get("cert1").strip(),
            "cert2_raw": get("cert2").strip(),
            "issues": [],
        }
        if not any((record["name"], record["email_raw"], record["cert1_raw"], record["cert2_raw"])):
            continue
        record["email_norm"] = record["email_raw"].casefold() if EMAIL_RE.fullmatch(record["email_raw"]) else None
        if not record["name"]:
            record["issues"].append("missing_name")
        if not record["email_raw"]:
            record["issues"].append("missing_email")
        elif record["email_norm"] is None:
            record["issues"].append("invalid_email")
        for key, issue in (("cert1_raw", "missing_certificate_1"), ("cert2_raw", "missing_certificate_2")):
            if not record[key]:
                record["issues"].append(issue)
        for key, issue in (("cert1_raw", "invalid_certificate_1_url"), ("cert2_raw", "invalid_certificate_2_url")):
            if record[key] and not safe_certificate_url(record[key]):
                record["issues"].append(issue)
        record["cert1_url"] = record["cert1_raw"] if record["cert1_raw"] and safe_certificate_url(record["cert1_raw"]) else None
        record["cert2_url"] = record["cert2_raw"] if record["cert2_raw"] and safe_certificate_url(record["cert2_raw"]) else None
        records.append(record)
    return records


def load_records(path: str | Path) -> list[dict[str, Any]]:
    source = Path(path).expanduser()
    if not source.is_file():
        raise FileNotFoundError(f"Input file does not exist: {source}")
    suffix = source.suffix.casefold()
    if suffix == ".csv":
        return _load_csv(source)
    if suffix in {".xlsx", ".xlsm"}:
        return _load_xlsx(source)
    raise ValueError("Supported import formats are CSV and XLSX")


def validate_records(records: list[dict[str, Any]]) -> dict[str, Any]:
    by_email: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        if record["email_norm"]:
            by_email[record["email_norm"]].append(record)
        if record["name"]:
            by_name[record["name"].casefold()].append(record)

    duplicate_email_groups = [group for group in by_email.values() if len(group) > 1]
    duplicate_name_groups = [group for group in by_name.values() if len(group) > 1]
    for group in duplicate_email_groups:
        for record in group:
            record["issues"].append("duplicate_email")
    for group in duplicate_name_groups:
        for record in group:
            record["issues"].append("duplicate_name_review")

    identity_valid = sum(
        1 for record in records
        if record["name"] and record["email_norm"] and len(by_email[record["email_norm"]]) == 1
    )
    issue_counts = Counter(issue for record in records for issue in set(record["issues"]))
    records_needing_correction = sum(bool(record["issues"]) for record in records)
    row_issues = [
        {"source_row": record["source_row"], "issues": sorted(set(record["issues"]))}
        for record in records if record["issues"]
    ]
    return {
        "total_participants": len(records),
        "valid_participant_records": identity_valid,
        "duplicate_email_ids": len(duplicate_email_groups),
        "duplicate_email_rows": sum(len(group) for group in duplicate_email_groups),
        "missing_email_ids": issue_counts["missing_email"],
        "invalid_email_ids": issue_counts["invalid_email"],
        "missing_names": issue_counts["missing_name"],
        "duplicate_participant_name_groups": len(duplicate_name_groups),
        "duplicate_participant_name_rows": sum(len(group) for group in duplicate_name_groups),
        "missing_certificate_1_links": issue_counts["missing_certificate_1"],
        "missing_certificate_2_links": issue_counts["missing_certificate_2"],
        "invalid_certificate_urls": issue_counts["invalid_certificate_1_url"] + issue_counts["invalid_certificate_2_url"],
        "records_requiring_correction": records_needing_correction,
        "row_issues": row_issues,
        "definition": "Valid participant records have a nonblank name and a valid, unique email. Missing or invalid certificate links are reported separately and do not discard a participant row.",
    }


def import_records(records: list[dict[str, Any]], db_path: str | Path | None = None) -> dict[str, Any]:
    report = validate_records(records)
    connection = connect(db_path)
    try:
        with connection:
            connection.execute("DELETE FROM participants")
            for record in records:
                email_norm = record["email_norm"]
                # Preserve every nonblank source row; duplicate emails remain ambiguous at lookup time.
                cert1_token = secrets.token_urlsafe(32) if record["cert1_url"] else None
                cert2_token = secrets.token_urlsafe(32) if record["cert2_url"] else None
                connection.execute(
                    """INSERT INTO participants
                       (id, source_row, name, email_norm, cert1_url, cert2_url, cert1_token, cert2_token,
                        cert1_token_hash, cert2_token_hash)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    (
                        uuid.uuid4().hex,
                        record["source_row"],
                        record["name"] or "Participant",
                        email_norm,
                        record["cert1_url"],
                        record["cert2_url"],
                        cert1_token,
                        cert2_token,
                        hashlib.sha256(cert1_token.encode()).hexdigest() if cert1_token else None,
                        hashlib.sha256(cert2_token.encode()).hexdigest() if cert2_token else None,
                    ),
                )
    finally:
        connection.close()
    return report


def _write_report(report: dict[str, Any], target: Path) -> None:
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix="validation-", suffix=".tmp", dir=target.parent)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as output:
            json.dump(report, output, ensure_ascii=False, indent=2)
            output.write("\n")
        os.replace(temporary, target)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate and import a private participant roster")
    parser.add_argument("source", help="CSV or XLSX roster")
    parser.add_argument("--database", default=str(database_path()), help="Private SQLite database path")
    parser.add_argument("--report", default=str(data_dir() / "validation-report.json"), help="Private JSON report path")
    args = parser.parse_args()
    try:
        records = load_records(args.source)
        report = import_records(records, args.database)
        _write_report(report, Path(args.report))
    except Exception as exc:
        print(f"Import failed: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2

    print("Participant roster validation and import completed")
    for key in (
        "total_participants", "valid_participant_records", "duplicate_email_ids",
        "missing_email_ids", "missing_names", "missing_certificate_1_links",
        "missing_certificate_2_links", "invalid_certificate_urls", "records_requiring_correction",
    ):
        print(f"{key}: {report[key]}")
    print(f"Private database: {args.database}")
    print(f"Private validation report: {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
